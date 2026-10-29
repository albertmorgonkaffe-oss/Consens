"""Avropa.se - Statens inköpscentral vid Kammarkollegiet (statliga ramavtal).

Källan är den officiella Excel-exporten "Alla giltiga ramavtal"
(/giltigaramavtal/excel) - stabilare än HTML och sanktionerad för
återanvändning (avropa.se har öppna data och publikt API med nyckel).

Kolumnerna läses TOLERANT: rubrikraden identifieras och kolumner mappas via
nyckelord (namn, giltigt till, leverantör...). Om mappningen misslyckas
loggas de faktiska rubrikerna så justeringen blir en enradsfix. Excel-filen
kan ha en rad per leverantör - rader grupperas per ramavtal och
leverantörerna samlas ihop.

Samma registeregenhet som Adda: förlängda giltighetstider uppdaterar
avtalet och rensar en inaktuell utgångssignal."""
import io
import json
import logging
import re
from datetime import date, datetime

from ..db import get_db
from ..matchning import bygg_namnindex, hamta_eller_skapa
from .bas import kor_avtalskalla
from .gemensamt import TIMEOUT, http_session

log = logging.getLogger(__name__)

EXCEL_URL = "https://www.avropa.se/giltigaramavtal/excel"
KOPARE = "Kammarkollegiet"
_DATUM = re.compile(r"(\d{4}-\d{2}-\d{2})")

# rubriknyckelord -> fält. Verifierat mot skarp export 2026-08-10:
# Avtalsnummer | Leverantör | Organisationsnummer | Ramavtalsområde |
# Delområde | Giltig från | Giltig till | Max förl. till
_KOLUMNER = {
    "omrade": ("ramavtalsområde", "avtalsområde", "ramavtal", "kategori"),
    "delomrade": ("delområde",),
    "slut": ("giltig till", "giltigt till", "slutdatum", "t.o.m",
             "till och med"),
    "start": ("giltig från", "giltigt från", "startdatum", "fr.o.m",
              "från och med"),
    "maxforl": ("max förl", "förläng"),
    "leverantor": ("leverantör",),
    "lank": ("länk", "url", "webbadress"),
}


def _till_datum(v) -> str | None:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    m = _DATUM.search(str(v or ""))
    return m.group(1) if m else None


def _slugga(text: str) -> str:
    t = (text or "").lower()
    for a, b in (("å", "a"), ("ä", "a"), ("ö", "o"), ("é", "e")):
        t = t.replace(a, b)
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return t[:120] or "ramavtal"


def tolka_arbetsbok(rader) -> list[dict]:
    """rader = iterable av tupler (cellvärden). Bruten ut för testbarhet.
    Returnerar grupperade ramavtal."""
    rader = [r for r in rader if any(c not in (None, "") for c in r)]
    if not rader:
        return []
    # Rubrikraden: första raden vars celler matchar minst namn + slut
    kol: dict = {}
    start_index = 0
    for i, rad in enumerate(rader[:10]):
        celler = [str(c or "").strip().lower() for c in rad]
        kandidat = {}
        for falt, nycklar in _KOLUMNER.items():
            traff = None
            # Exakta rubrikträffar prioriteras före delsträngar så att
            # 'delområde' inte kapas av 'område'-nycklarna.
            for exakt_pass in (True, False):
                for j, cell in enumerate(celler):
                    if falt == "omrade" and "delområde" in cell:
                        continue
                    if any((cell == n) if exakt_pass else (n in cell)
                           for n in nycklar):
                        traff = j
                        break
                if traff is not None:
                    break
            if traff is not None:
                kandidat[falt] = traff
        if "omrade" in kandidat and "slut" in kandidat:
            kol, start_index = kandidat, i + 1
            break
    if not kol:
        log.error("avropa: hittade ingen rubrikrad. Första raderna: %s",
                  [r[:8] for r in rader[:3]])
        return []

    # En rad per leverantörskontrakt (t.o.m. per hotellzon) - grupperas
    # per ramavtalsOMRÅDE så flödet får ~100 strategiska avtal i stället
    # för tusentals zonrader. Slutdatum = senaste "Giltig till" i området.
    grupper: dict[str, dict] = {}
    for rad in rader[start_index:]:
        def _c(falt):
            j = kol.get(falt)
            return (str(rad[j]).strip() if j is not None and j < len(rad)
                    and rad[j] not in (None, "") else "")
        namn = _c("omrade")
        slut = _till_datum(rad[kol["slut"]] if kol["slut"] < len(rad) else None)
        if not namn or not slut:
            continue
        g = grupper.setdefault(_slugga(namn), dict(
            namn=namn, slutdatum=slut, startdatum=None, maxforl=None,
            lank=_c("lank"), leverantorer=[], delomraden=set()))
        g["slutdatum"] = max(g["slutdatum"], slut)
        start = (_till_datum(rad[kol["start"]])
                 if "start" in kol and kol["start"] < len(rad) else None)
        if start and (g["startdatum"] is None or start < g["startdatum"]):
            g["startdatum"] = start
        maxforl = (_till_datum(rad[kol["maxforl"]])
                   if "maxforl" in kol and kol["maxforl"] < len(rad) else None)
        if maxforl and (g["maxforl"] is None or maxforl > g["maxforl"]):
            g["maxforl"] = maxforl
        if (d := _c("delomrade")):
            g["delomraden"].add(d)
        lev = _c("leverantor")
        if lev and lev not in g["leverantorer"]:
            g["leverantorer"].append(lev)
    return list(grupper.values())


def hamta():
    from openpyxl import load_workbook
    db = get_db()
    session = http_session()
    resp = session.get(EXCEL_URL, timeout=TIMEOUT)
    resp.raise_for_status()
    wb = load_workbook(io.BytesIO(resp.content), read_only=True,
                       data_only=True)
    rader = [tuple(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    exakt, normaliserat, per_orgnr = bygg_namnindex(db)

    for g in tolka_arbetsbok(rader):
        extern_id = f"avropa:{_slugga(g['namn'])}"
        kand = db.execute("SELECT id, slutdatum FROM avtal WHERE extern_id=?",
                          (extern_id,)).fetchone()
        if kand:
            if kand["slutdatum"] != g["slutdatum"]:
                db.execute("UPDATE avtal SET slutdatum=? WHERE id=?",
                           (g["slutdatum"], kand["id"]))
                db.execute("DELETE FROM signal WHERE extern_id=?",
                           (f"utgang:{extern_id}",))
                db.commit()
            continue

        org, _, _ = hamta_eller_skapa(db, KOPARE, None,
                                      exakt, normaliserat, per_orgnr)
        if org is None:
            continue
        lev = g["leverantorer"]
        levtext = (", ".join(lev) if len(lev) <= 6 else
                   ", ".join(lev[:6]) + f" med flera ({len(lev)} leverantörer)")
        detaljer = {k: v for k, v in {
            "Avtalsform": "Statligt ramavtal (Statens inköpscentral)",
            "Kan som längst förlängas till": g["maxforl"],
            "Delområden": (str(len(g["delomraden"]))
                           if len(g["delomraden"]) > 1 else None),
        }.items() if v}
        yield dict(
            organisation_id=org["id"],
            titel=g["namn"][:500],
            leverantor=levtext[:500] or None,
            startdatum=g["startdatum"],
            slutdatum=g["slutdatum"],
            kalla="avropa",
            extern_id=extern_id,
            url=g["lank"] or "https://www.avropa.se/ramavtal/",
            detaljer=json.dumps(detaljer, ensure_ascii=False),
        )


def kor(torr=False):
    return kor_avtalskalla("avropa", hamta, torr=torr)
